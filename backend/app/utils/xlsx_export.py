import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from typing import List


def export_xlsx(questions, responses) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")

    headers = ["Response ID", "Submitted At", "IP"] + [q.label or f"Question {q.id}" for q in questions]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, resp in enumerate(responses, 2):
        ws.cell(row=row_idx, column=1, value=resp.id)
        ws.cell(row=row_idx, column=2, value=resp.submitted_at.isoformat())
        ws.cell(row=row_idx, column=3, value=resp.ip or "")
        for col_idx, q in enumerate(questions, 4):
            ans = next((a for a in resp.answers if a.question_id == q.id), None)
            if ans:
                if ans.value_json:
                    val = json.loads(ans.value_json)
                    ws.cell(row=row_idx, column=col_idx, value=json.dumps(val) if isinstance(val, (list, dict)) else str(val))
                else:
                    ws.cell(row=row_idx, column=col_idx, value=ans.value_text or "")

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
